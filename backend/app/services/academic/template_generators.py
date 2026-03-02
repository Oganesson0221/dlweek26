import os
from datetime import datetime
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

from app.core.config import settings
from app.db.mongodb import save_generated_doc


async def generate_word_doc(assignment: dict) -> dict:
    """Generate a Word document template for an assignment."""
    doc = Document()
    
    # Title
    doc.add_heading(assignment.get("title", "Assignment"), 0)
    
    # Course info
    doc.add_paragraph(f"Course: {assignment.get('course_code', 'N/A')}")
    doc.add_paragraph(f"Due Date: {assignment.get('due_at', 'N/A')}")
    doc.add_paragraph(f"Weight: {assignment.get('weight', 0) * 100:.0f}%")
    
    doc.add_heading("Description", level=1)
    doc.add_paragraph(assignment.get("description", "No description provided."))
    
    doc.add_heading("Your Work", level=1)
    doc.add_paragraph("[Start writing your assignment here]")
    
    doc.add_heading("References", level=1)
    doc.add_paragraph("[Add your references here]")
    
    # Save file
    os.makedirs(settings.GENERATED_DIR, exist_ok=True)
    file_name = f"{assignment.get('course_code', 'course')}_{assignment.get('title', 'assignment').replace(' ', '_')}.docx"
    file_path = os.path.join(settings.GENERATED_DIR, file_name)
    doc.save(file_path)
    
    # Save to MongoDB
    saved = save_generated_doc(
        assignment_id=assignment["id"],
        doc_type="word",
        file_name=file_name,
        file_path=file_path
    )
    
    return {
        "id": saved["id"],
        "file_name": file_name,
        "doc_type": "word",
        "download_url": f"/academic/templates/download/{saved['id']}"
    }


async def generate_ppt_doc(assignment: dict) -> dict:
    """Generate a PowerPoint template for an assignment."""
    prs = Presentation()
    
    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = assignment.get("title", "Presentation")
    subtitle.text = f"Course: {assignment.get('course_code', 'N/A')}\nDue: {assignment.get('due_at', 'N/A')}"
    
    # Content slide
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "Overview"
    tf = body_shape.text_frame
    tf.text = assignment.get("description", "Add your content here")
    
    # Add more slides
    for i in range(3):
        slide = prs.slides.add_slide(bullet_slide_layout)
        title_shape = slide.shapes.title
        title_shape.text = f"Section {i + 1}"
        body_shape = slide.shapes.placeholders[1]
        body_shape.text_frame.text = "Add your content here"
    
    # Conclusion slide
    slide = prs.slides.add_slide(bullet_slide_layout)
    slide.shapes.title.text = "Conclusion"
    slide.shapes.placeholders[1].text_frame.text = "Summary and key takeaways"
    
    # Save file
    os.makedirs(settings.GENERATED_DIR, exist_ok=True)
    file_name = f"{assignment.get('course_code', 'course')}_{assignment.get('title', 'presentation').replace(' ', '_')}.pptx"
    file_path = os.path.join(settings.GENERATED_DIR, file_name)
    prs.save(file_path)
    
    # Save to MongoDB
    saved = save_generated_doc(
        assignment_id=assignment["id"],
        doc_type="ppt",
        file_name=file_name,
        file_path=file_path
    )
    
    return {
        "id": saved["id"],
        "file_name": file_name,
        "doc_type": "ppt",
        "download_url": f"/academic/templates/download/{saved['id']}"
    }
