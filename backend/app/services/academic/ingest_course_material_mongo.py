import os
import re
import json
from datetime import datetime
from typing import Dict, Any, List

from fastapi import UploadFile, HTTPException
from app.core.config import settings
from app.db.mongodb import upsert_outline, replace_components, add_topic


# -------------------------
# File save + text extract
# -------------------------

async def save_upload(file: UploadFile) -> str:
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    safe_name = re.sub(r"[^a-zA-Z0-9._-]", "_", file.filename or "upload.bin")
    path = os.path.join(settings.UPLOAD_DIR, safe_name)

    contents = await file.read()
    with open(path, "wb") as f:
        f.write(contents)

    return path


def extract_text_basic(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"PDF extract failed: {e}")

    if ext in [".ppt", ".pptx"]:
        try:
            from pptx import Presentation
            prs = Presentation(path)
            lines: List[str] = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        lines.append(shape.text)
            return "\n".join(lines)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"PPT extract failed: {e}")

    if ext == ".docx":
        try:
            import docx
            d = docx.Document(path)
            return "\n".join(p.text for p in d.paragraphs if p.text)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"DOCX extract failed: {e}")

    raise HTTPException(status_code=400, detail="Unsupported file type. Upload PDF/PPTX/DOCX.")


def clean_text(text: str, max_chars: int = 160000) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > max_chars:
        text = text[:max_chars]
    return text


def chunk_text(text: str, max_chars: int = 12000, overlap: int = 600) -> List[str]:
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: List[str] = []
    cur: List[str] = []
    cur_len = 0

    def flush():
        nonlocal cur, cur_len
        if not cur:
            return
        chunk = "\n\n".join(cur).strip()
        if chunk:
            chunks.append(chunk)
        cur = []
        cur_len = 0

    for p in paras:
        add_len = len(p) + 2
        if cur_len + add_len > max_chars and cur:
            flush()
        cur.append(p)
        cur_len += add_len

    flush()

    if overlap > 0 and len(chunks) > 1:
        out: List[str] = []
        for i, c in enumerate(chunks):
            if i == 0:
                out.append(c)
            else:
                prev_tail = chunks[i - 1][-overlap:]
                out.append(prev_tail + "\n\n" + c)
        chunks = out

    return chunks


# -------------------------
# OpenAI schemas
# -------------------------

CHUNK_SCHEMA: Dict[str, Any] = {
    "name": "course_chunk_extraction",
    "schema": {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "instructor_candidates": {"type": "array", "items": {"type": "string"}},
            "outline_bullets": {"type": "array", "items": {"type": "string"}},
            "components": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "name": {"type": "string"},
                        "weight": {"type": ["number", "null"], "minimum": 0, "maximum": 1}
                    },
                    "required": ["name", "weight"]
                }
            },
            "topics": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "title": {"type": "string"},
                        "parent_title": {"type": ["string", "null"]},
                        "order_index": {"type": "integer", "minimum": 1}
                    },
                    "required": ["title", "parent_title", "order_index"]
                }
            }
        },
        "required": ["instructor_candidates", "outline_bullets", "components", "topics"]
    }
}

FINAL_SCHEMA: Dict[str, Any] = {
    "name": "course_extraction_final",
    "schema": {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "outline_description": {"type": "string"},
            "instructor": {"type": "string"},
            "components": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "name": {"type": "string"},
                        "weight": {"type": "number", "minimum": 0, "maximum": 1}
                    },
                    "required": ["name", "weight"]
                }
            },
            "topics": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "title": {"type": "string"},
                        "parent_title": {"type": ["string", "null"]},
                        "order_index": {"type": "integer", "minimum": 1}
                    },
                    "required": ["title", "parent_title", "order_index"]
                }
            }
        },
        "required": ["outline_description", "instructor", "components", "topics"]
    }
}


def _openai_client():
    api_key = settings.OPENAI_API_KEY
    if not api_key:
        return None
    from openai import OpenAI
    return OpenAI(api_key=api_key)


def _openai_model() -> str:
    return os.getenv("DEFAULT_MODEL", "gpt-4.1")


def _openai_temp() -> float:
    try:
        return float(os.getenv("TEMPERATURE", "0.2"))
    except Exception:
        return 0.2


def extract_from_chunk_openai(course_code: str, chunk: str) -> Dict[str, Any]:
    client = _openai_client()
    if client is None:
        raise HTTPException(status_code=500, detail="OPENAI_API_KEY not set")

    resp = client.responses.create(
        model=_openai_model(),
        input=[
            {"role": "system", "content": (
                "You extract structure from a course document chunk. "
                "Find assessment components (grading), major topics and subtopics, and outline bullets. "
                "If weights are not explicitly stated, set weight to null. "
                "Return only JSON matching the schema."
            )},
            {"role": "user", "content": (
                f"Course: {course_code}\n\n"
                "Extract from this chunk:\n"
                "- instructor_candidates: any instructor names found\n"
                "- outline_bullets: short bullets describing what the course covers\n"
                "- components: assessments/grading items (name + weight if specified)\n"
                "- topics: main topics and subtopics (use parent_title to form hierarchy)\n\n"
                f"CHUNK:\n{chunk}"
            )}
        ],
        text={"format": {"type": "json_schema", "name": CHUNK_SCHEMA["name"], "schema": CHUNK_SCHEMA["schema"]}},
        temperature=_openai_temp(),
    )

    return json.loads(resp.output_text)


def merge_signals_openai(course_code: str, signals: List[Dict[str, Any]]) -> Dict[str, Any]:
    client = _openai_client()
    if client is None:
        raise HTTPException(status_code=500, detail="OPENAI_API_KEY not set")

    compact = json.dumps(signals, ensure_ascii=False)[:250000]

    resp = client.responses.create(
        model=_openai_model(),
        input=[
            {"role": "system", "content": (
                "You consolidate multiple chunk extractions into one clean course structure. "
                "Prefer explicit weights. Deduplicate similar components. "
                "If weights are missing, infer reasonable weights but keep them realistic, "
                "and make weights sum to 1.0."
                "Return only JSON matching the schema."
            )},
            {"role": "user", "content": (
                f"Course: {course_code}\n\n"
                "Given these chunk-level extractions, produce FINAL course structure.\n"
                "- outline_description: coherent paragraph(s)\n"
                "- instructor: best guess (or empty)\n"
                "- components: assessments with weights summing to 1.0\n"
                "- topics: deduped hierarchy with order_index increasing\n\n"
                f"CHUNK_SIGNALS_JSON:\n{compact}"
            )}
        ],
        text={"format": {"type": "json_schema", "name": FINAL_SCHEMA["name"], "schema": FINAL_SCHEMA["schema"]}},
        temperature=_openai_temp(),
    )

    return json.loads(resp.output_text)


def persist_topics(course_code: str, topics: List[Dict[str, Any]]) -> List[Any]:
    """Two-pass insert using parent_title."""
    created: List[Any] = []
    title_to_id: Dict[str, str] = {}

    roots = [t for t in topics if t.get("parent_title") in [None, ""]]
    kids = [t for t in topics if t.get("parent_title") not in [None, ""]]

    for t in sorted(roots, key=lambda x: x.get("order_index", 1)):
        row = add_topic(course_code, None, t["title"], int(t.get("order_index", 1)))
        created.append(row)
        title_to_id[t["title"].strip().lower()] = row["id"]

    for t in sorted(kids, key=lambda x: x.get("order_index", 1)):
        parent_title = (t.get("parent_title") or "").strip().lower()
        parent_id = title_to_id.get(parent_title)
        row = add_topic(course_code, parent_id, t["title"], int(t.get("order_index", 1)))
        created.append(row)
        title_to_id[t["title"].strip().lower()] = row["id"]

    return created


async def ingest_course_material(course_id: str, course_code: str, file: UploadFile) -> Dict[str, Any]:
    path = await save_upload(file)
    raw = extract_text_basic(path)
    text = clean_text(raw)

    chunks = chunk_text(text, max_chars=12000, overlap=600)

    signals: List[Dict[str, Any]] = []
    for i, ch in enumerate(chunks):
        try:
            sig = extract_from_chunk_openai(course_code, ch)
            sig["_chunk_index"] = i + 1
            signals.append(sig)
        except Exception as e:
            signals.append({
                "_chunk_index": i + 1,
                "_error": str(e),
                "instructor_candidates": [],
                "outline_bullets": [],
                "components": [],
                "topics": [],
            })

    final = merge_signals_openai(course_code, signals)

    # Persist to MongoDB
    upsert_outline(course_code, final["outline_description"], final.get("instructor", ""))

    comps = [(c["name"], float(c["weight"])) for c in final["components"]]
    replace_components(course_code, comps)

    created_topics = persist_topics(course_code, final["topics"])

    return {
        "upload_path": path,
        "file_name": file.filename,
        "text_chars": len(text),
        "chunks": len(chunks),
        "outline_preview": final["outline_description"][:800],
        "components_extracted": final["components"],
        "topics_extracted": final["topics"],
        "components_saved": len(comps),
        "topics_saved": len(created_topics),
        "saved_topic_ids": [t["id"] for t in created_topics],
        "timestamp": datetime.utcnow().isoformat(),
        "debug_chunk_signals_sample": signals[:2],
    }
