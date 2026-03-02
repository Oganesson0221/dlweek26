from pathlib import Path
from datetime import datetime
from sqlmodel import Session, select
from fastapi import HTTPException

from app.core.config import settings
from app.models.academic import (
    Assignment,
    Course,
    CourseOutline,
    CourseComponent,
    TopicNode,
    GeneratedDocument,
)

def generate_word(session: Session, assignment_id: int):
    # 1) Find assignment
    a = session.exec(select(Assignment).where(Assignment.id == assignment_id)).first()
    if not a:
        raise HTTPException(status_code=404, detail="Assignment not found")

    # 2) Find course + related data
    course = session.exec(select(Course).where(Course.id == a.course_id)).first()
    if not course:
        raise HTTPException(status_code=404, detail="Course not found for assignment")

    outline = session.exec(select(CourseOutline).where(CourseOutline.course_id == course.id)).first()
    components = session.exec(select(CourseComponent).where(CourseComponent.course_id == course.id)).all()
    topics = session.exec(
        select(TopicNode).where(TopicNode.course_id == course.id).order_by(TopicNode.order_index)
    ).all()

    # 3) Build Word doc
    try:
        from docx import Document
    except Exception:
        raise HTTPException(status_code=500, detail="python-docx not installed. pip install python-docx")

    out_dir = Path(settings.GENERATED_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)

    safe_course = course.code.replace("/", "_").replace(" ", "_")
    safe_assign = a.title.replace("/", "_").replace(" ", "_")
    file_name = f"{safe_course}_{safe_assign}_course_summary.docx"
    file_path = out_dir / file_name

    doc = Document()
    doc.add_heading(f"{course.code} - {course.name}", level=1)
    doc.add_paragraph(f"Term: {course.term}")
    doc.add_paragraph(f"Generated: {datetime.utcnow().isoformat()}")

    doc.add_heading("Assignment", level=2)
    doc.add_paragraph(f"Title: {a.title}")
    doc.add_paragraph(f"Due: {a.due_at.isoformat()}")
    doc.add_paragraph(f"Weight: {a.weight}")

    doc.add_heading("Course Outline", level=2)
    if outline and outline.description:
        doc.add_paragraph(outline.description)
    else:
        doc.add_paragraph("No outline found in DB yet.")

    doc.add_heading("Assessment Components", level=2)
    if components:
        for c in components:
            pct = round(c.weight * 100, 1)
            doc.add_paragraph(f"- {c.name}: {pct}%")
    else:
        doc.add_paragraph("No components found in DB yet.")

    doc.add_heading("Topics", level=2)
    if topics:
        for t in topics:
            doc.add_paragraph(f"{t.order_index}. {t.title}")
    else:
        doc.add_paragraph("No topics found in DB yet.")

    doc.save(str(file_path))

    # 4) Save generated doc record
    gd = GeneratedDocument(
        assignment_id=a.id,
        doc_type="word",
        file_name=file_name,
        file_path=str(file_path),
    )
    session.add(gd)
    session.commit()
    session.refresh(gd)

    return {
        "doc_id": gd.id,
        "doc_type": gd.doc_type,
        "file_name": gd.file_name,
        "download_path": f"/academic/templates/download/{gd.id}",
        "course_code": course.code,
        "assignment_id": a.id,
    }