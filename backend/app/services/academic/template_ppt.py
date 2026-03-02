from pathlib import Path
from datetime import datetime
from sqlmodel import Session, select
from fastapi import HTTPException

from app.core.config import settings
from app.models.academic import (
    Assignment,
    Course,
    CourseOutline,
    CourseComponent,
    TopicNode,
    GeneratedDocument,
)

def generate_ppt(session: Session, assignment_id: int):
    a = session.exec(select(Assignment).where(Assignment.id == assignment_id)).first()
    if not a:
        raise HTTPException(status_code=404, detail="Assignment not found")

    course = session.exec(select(Course).where(Course.id == a.course_id)).first()
    if not course:
        raise HTTPException(status_code=404, detail="Course not found for assignment")

    outline = session.exec(select(CourseOutline).where(CourseOutline.course_id == course.id)).first()
    components = session.exec(select(CourseComponent).where(CourseComponent.course_id == course.id)).all()
    topics = session.exec(
        select(TopicNode).where(TopicNode.course_id == course.id).order_by(TopicNode.order_index)
    ).all()

    try:
        from pptx import Presentation
    except Exception:
        raise HTTPException(status_code=500, detail="python-pptx not installed. pip install python-pptx")

    out_dir = Path(settings.GENERATED_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)

    safe_course = course.code.replace("/", "_").replace(" ", "_")
    safe_assign = a.title.replace("/", "_").replace(" ", "_")
    file_name = f"{safe_course}_{safe_assign}_course_summary.pptx"
    file_path = out_dir / file_name

    prs = Presentation()

    # Slide 1: Overview
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    slide.shapes.title.text = f"{course.code} - {course.name}"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    body.text = f"Term: {course.term}"
    body.add_paragraph().text = f"Assignment: {a.title}"
    body.add_paragraph().text = f"Due: {a.due_at.isoformat()}"
    body.add_paragraph().text = f"Generated: {datetime.utcnow().isoformat()}"

    # Slide 2: Outline
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Course Outline"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    if outline and outline.description:
        body.text = outline.description[:1200]
        if len(outline.description) > 1200:
            body.add_paragraph().text = "(Outline truncated for slide)"
    else:
        body.text = "No outline found in DB yet."

    # Slide 3: Components + Topics
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Assessments and Topics"
    body = slide.shapes.placeholders[1].text_frame
    body.clear()

    body.text = "Assessment Components:"
    if components:
        for c in components:
            pct = round(c.weight * 100, 1)
            body.add_paragraph().text = f"- {c.name}: {pct}%"
    else:
        body.add_paragraph().text = "- No components found"

    body.add_paragraph().text = ""
    body.add_paragraph().text = "Topics:"
    if topics:
        for t in topics[:20]:
            body.add_paragraph().text = f"{t.order_index}. {t.title}"
        if len(topics) > 20:
            body.add_paragraph().text = "(Topics truncated)"
    else:
        body.add_paragraph().text = "- No topics found"

    prs.save(str(file_path))

    gd = GeneratedDocument(
        assignment_id=a.id,
        doc_type="ppt",
        file_name=file_name,
        file_path=str(file_path),
    )
    session.add(gd)
    session.commit()
    session.refresh(gd)

    return {
        "doc_id": gd.id,
        "doc_type": gd.doc_type,
        "file_name": gd.file_name,
        "download_path": f"/academic/templates/download/{gd.id}",
        "course_code": course.code,
        "assignment_id": a.id,
    }