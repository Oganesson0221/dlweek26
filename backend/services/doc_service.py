from docx import Document
from io import BytesIO

def generate_assignment_template(course_code: str, assignment_title: str, user_name: str) -> BytesIO:
    """Creates a basic Word document template for an assignment."""
    doc = Document()
    doc.add_heading(f'{course_code} - {assignment_title}', 0)
    
    doc.add_paragraph(f'Student Name: {user_name}')
    doc.add_paragraph('Date: _______________')
    doc.add_heading('Introduction', level=1)
    doc.add_paragraph('Start typing your introduction here...')
    
    doc.add_heading('Main Content', level=1)
    doc.add_paragraph('Your main points go here...')
    
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream